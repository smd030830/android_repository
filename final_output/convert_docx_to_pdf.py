import sys
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


def register_fonts():
    font = r"C:\Windows\Fonts\malgun.ttf"
    font_bold = r"C:\Windows\Fonts\malgunbd.ttf"
    pdfmetrics.registerFont(TTFont("Malgun", font))
    pdfmetrics.registerFont(TTFont("MalgunBold", font_bold))


def build_styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="KBody", parent=styles["Normal"], fontName="Malgun", fontSize=9.3, leading=13, spaceAfter=5))
    styles.add(ParagraphStyle(name="KTitle", parent=styles["Title"], fontName="MalgunBold", fontSize=22, leading=27, textColor=colors.HexColor("#0B2545"), alignment=1, spaceAfter=14))
    styles.add(ParagraphStyle(name="KH1", parent=styles["Heading1"], fontName="MalgunBold", fontSize=15, leading=19, textColor=colors.HexColor("#2E74B5"), spaceBefore=10, spaceAfter=7))
    styles.add(ParagraphStyle(name="KH2", parent=styles["Heading2"], fontName="MalgunBold", fontSize=12, leading=15, textColor=colors.HexColor("#2E74B5"), spaceBefore=8, spaceAfter=5))
    styles.add(ParagraphStyle(name="KH3", parent=styles["Heading3"], fontName="MalgunBold", fontSize=10.5, leading=13, textColor=colors.HexColor("#1F4D78"), spaceBefore=6, spaceAfter=4))
    styles.add(ParagraphStyle(name="KCell", parent=styles["Normal"], fontName="Malgun", fontSize=7.6, leading=10, wordWrap="CJK"))
    styles.add(ParagraphStyle(name="KCellHead", parent=styles["Normal"], fontName="MalgunBold", fontSize=7.8, leading=10, wordWrap="CJK", textColor=colors.HexColor("#0B2545")))
    styles.add(ParagraphStyle(name="KCode", parent=styles["Code"], fontName="Courier", fontSize=7.2, leading=8.8, backColor=colors.HexColor("#F4F6F9"), borderPadding=4, spaceAfter=5))
    return styles


def text_of_para(para):
    return "".join(run.text for run in para.runs).strip()


def para_flow(text, style):
    if not text:
        return Spacer(1, 4)
    return Paragraph(escape(text).replace("\n", "<br/>"), style)


def is_page_break(para):
    return bool(para._element.xpath('.//w:br[@w:type="page"]'))


def docx_to_pdf(input_path, output_path):
    register_fonts()
    styles = build_styles()
    doc = Document(input_path)
    story = []
    usable = 6.5 * inch

    for child in doc.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            para = next((p for p in doc.paragraphs if p._p is child), None)
            if para is None:
                continue
            text = text_of_para(para)
            if is_page_break(para):
                if text:
                    story.append(para_flow(text, styles["KBody"]))
                story.append(PageBreak())
                continue
            style_name = para.style.name if para.style is not None else ""
            if not text:
                story.append(Spacer(1, 4))
            elif text.startswith("DogLog\n"):
                story.append(para_flow(text, styles["KTitle"]))
            elif style_name == "Heading 1":
                story.append(para_flow(text, styles["KH1"]))
            elif style_name == "Heading 2":
                story.append(para_flow(text, styles["KH2"]))
            elif style_name == "Heading 3":
                story.append(para_flow(text, styles["KH3"]))
            elif any(run.font.name == "Consolas" for run in para.runs):
                story.append(para_flow(text, styles["KCode"]))
            elif style_name in ("List Bullet", "List Number"):
                marker = "• " if style_name == "List Bullet" else "· "
                story.append(para_flow(marker + text, styles["KBody"]))
            else:
                story.append(para_flow(text, styles["KBody"]))
        elif tag == "tbl":
            tbl = next((t for t in doc.tables if t._tbl is child), None)
            if tbl is None:
                continue
            rows = []
            for row_index, row in enumerate(tbl.rows):
                cells = []
                for cell in row.cells:
                    cell_text = "\n".join(p.text for p in cell.paragraphs).strip()
                    cell_style = styles["KCellHead"] if row_index == 0 else styles["KCell"]
                    cells.append(Paragraph(escape(cell_text).replace("\n", "<br/>"), cell_style))
                rows.append(cells)
            if rows:
                col_count = len(rows[0])
                if col_count == 2:
                    widths = [1.45 * inch, usable - 1.45 * inch]
                elif col_count == 3:
                    widths = [1.25 * inch, 2.25 * inch, usable - 3.5 * inch]
                else:
                    widths = [usable / col_count] * col_count
                rt = Table(rows, colWidths=widths, repeatRows=1, splitByRow=1)
                rt.setStyle(TableStyle([
                    ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#C9D2DC")),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F2F4F7")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]))
                story.append(rt)
                story.append(Spacer(1, 6))

    pdf = SimpleDocTemplate(str(output_path), pagesize=letter, rightMargin=inch, leftMargin=inch, topMargin=inch, bottomMargin=inch)
    pdf.build(story)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("usage: convert_docx_to_pdf.py input.docx output.pdf")
    docx_to_pdf(Path(sys.argv[1]), Path(sys.argv[2]))
