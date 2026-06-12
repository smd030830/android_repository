from pathlib import Path
from xml.sax.saxutils import escape
from docx import Document
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

OUT = Path('final_output')
DOCX = OUT / '최종보고서_DogLog_DB보완.docx'
PDF = OUT / '최종보고서_DogLog_DB보완.pdf'
FONT = r'C:\Windows\Fonts\malgun.ttf'
FONT_BOLD = r'C:\Windows\Fonts\malgunbd.ttf'
pdfmetrics.registerFont(TTFont('Malgun', FONT))
pdfmetrics.registerFont(TTFont('MalgunBold', FONT_BOLD))

styles = getSampleStyleSheet()
styles.add(ParagraphStyle(name='KBody', parent=styles['Normal'], fontName='Malgun', fontSize=9.3, leading=13, spaceAfter=5))
styles.add(ParagraphStyle(name='KTitle', parent=styles['Title'], fontName='MalgunBold', fontSize=22, leading=27, textColor=colors.HexColor('#0B2545'), alignment=1, spaceAfter=14))
styles.add(ParagraphStyle(name='KH1', parent=styles['Heading1'], fontName='MalgunBold', fontSize=15, leading=19, textColor=colors.HexColor('#2E74B5'), spaceBefore=10, spaceAfter=7))
styles.add(ParagraphStyle(name='KH2', parent=styles['Heading2'], fontName='MalgunBold', fontSize=12, leading=15, textColor=colors.HexColor('#2E74B5'), spaceBefore=8, spaceAfter=5))
styles.add(ParagraphStyle(name='KH3', parent=styles['Heading3'], fontName='MalgunBold', fontSize=10.5, leading=13, textColor=colors.HexColor('#1F4D78'), spaceBefore=6, spaceAfter=4))
styles.add(ParagraphStyle(name='KCell', parent=styles['Normal'], fontName='Malgun', fontSize=7.6, leading=10, wordWrap='CJK'))
styles.add(ParagraphStyle(name='KCellHead', parent=styles['Normal'], fontName='MalgunBold', fontSize=7.8, leading=10, wordWrap='CJK', textColor=colors.HexColor('#0B2545')))
styles.add(ParagraphStyle(name='KCode', parent=styles['Code'], fontName='Courier', fontSize=7.2, leading=8.8, backColor=colors.HexColor('#F4F6F9'), borderPadding=4, spaceAfter=5))

story = []
doc = Document(DOCX)
usable = 6.5 * inch

def text_of_para(p):
    return ''.join(run.text for run in p.runs).strip()

def para_flow(text, style):
    if not text:
        return Spacer(1, 4)
    safe = escape(text).replace('\n', '<br/>')
    return Paragraph(safe, style)

def is_page_break(p):
    return bool(p._element.xpath('.//w:br[@w:type="page"]'))

for child in doc.element.body.iterchildren():
    tag = child.tag.split('}')[-1]
    if tag == 'p':
        # Find corresponding python-docx paragraph by element identity.
        dp = None
        for para in doc.paragraphs:
            if para._p is child:
                dp = para
                break
        if dp is None:
            continue
        txt = text_of_para(dp)
        if is_page_break(dp):
            if txt:
                story.append(para_flow(txt, styles['KBody']))
            story.append(PageBreak())
            continue
        sname = dp.style.name if dp.style is not None else ''
        if not txt:
            story.append(Spacer(1, 4))
        elif txt.startswith('DogLog\n'):
            story.append(para_flow(txt, styles['KTitle']))
        elif sname == 'Heading 1':
            story.append(para_flow(txt, styles['KH1']))
        elif sname == 'Heading 2':
            story.append(para_flow(txt, styles['KH2']))
        elif sname == 'Heading 3':
            story.append(para_flow(txt, styles['KH3']))
        elif any(run.font.name == 'Consolas' for run in dp.runs):
            story.append(para_flow(txt, styles['KCode']))
        elif sname in ('List Bullet', 'List Number'):
            marker = '• ' if sname == 'List Bullet' else '· '
            story.append(para_flow(marker + txt, styles['KBody']))
        else:
            story.append(para_flow(txt, styles['KBody']))
    elif tag == 'tbl':
        # Match table by element identity.
        dt = None
        for tbl in doc.tables:
            if tbl._tbl is child:
                dt = tbl
                break
        if dt is None:
            continue
        rows = []
        for r_idx, row in enumerate(dt.rows):
            cells = []
            for cell in row.cells:
                ctext = '\n'.join(p.text for p in cell.paragraphs).strip()
                cstyle = styles['KCellHead'] if r_idx == 0 else styles['KCell']
                cells.append(Paragraph(escape(ctext).replace('\n', '<br/>'), cstyle))
            rows.append(cells)
        if rows:
            n = len(rows[0])
            if n == 2:
                widths = [1.45 * inch, usable - 1.45 * inch]
            elif n == 3:
                widths = [1.25 * inch, 2.25 * inch, usable - 3.5 * inch]
            else:
                widths = [usable / n] * n
            rt = Table(rows, colWidths=widths, repeatRows=1, splitByRow=1)
            rt.setStyle(TableStyle([
                ('GRID', (0,0), (-1,-1), 0.35, colors.HexColor('#C9D2DC')),
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#F2F4F7')),
                ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
                ('LEFTPADDING', (0,0), (-1,-1), 5),
                ('RIGHTPADDING', (0,0), (-1,-1), 5),
                ('TOPPADDING', (0,0), (-1,-1), 4),
                ('BOTTOMPADDING', (0,0), (-1,-1), 4),
            ]))
            story.append(rt)
            story.append(Spacer(1, 6))

pdf = SimpleDocTemplate(str(PDF), pagesize=letter, rightMargin=inch, leftMargin=inch, topMargin=inch, bottomMargin=inch)
pdf.build(story)
print(PDF.resolve())
