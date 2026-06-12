from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont


BASE_DIR = Path(r"D:\github\android_repository\final_output\PJ산출물_DogLog_DB설계반영")
SOURCE_DOCX = BASE_DIR / "최종보고서_DogLog.docx"
OUTPUT_DOCX = BASE_DIR / "최종보고서_DogLog_사진포함.docx"
IMAGE_DIR = BASE_DIR / "db_capture_images"


def font(size, bold=False):
    path = r"C:\Windows\Fonts\malgunbd.ttf" if bold else r"C:\Windows\Fonts\malgun.ttf"
    return ImageFont.truetype(path, size)


def ellipsize(draw, text, fnt, max_width):
    text = str(text)
    if draw.textlength(text, font=fnt) <= max_width:
        return text
    while text and draw.textlength(text + "...", font=fnt) > max_width:
        text = text[:-1]
    return text + "..."


def draw_db_table(name, headers, rows, widths, output):
    row_h = 34
    title_h = 28
    pad_x = 10
    width = sum(widths) + 2
    height = title_h + row_h * (len(rows) + 1) + 2
    img = Image.new("RGB", (width, height), "#2f3437")
    draw = ImageDraw.Draw(img)
    header_font = font(14, True)
    cell_font = font(14, False)
    title_font = font(16, True)

    draw.rectangle([0, 0, width, title_h], fill="#252b2f")
    draw.text((8, 4), name, fill="#1ed760", font=title_font)
    draw.text((120, 5), "Enter a SQL expression to filter results (use Ctrl+Space)", fill="#c0c8ca", font=font(12))

    y = title_h
    x = 1
    for i, header in enumerate(headers):
        draw.rectangle([x, y, x + widths[i], y + row_h], fill="#4b5968", outline="#1f2428")
        draw.text((x + pad_x, y + 8), header, fill="#d9edf7", font=header_font)
        x += widths[i]

    y += row_h
    for r_idx, row in enumerate(rows):
        fill = "#3b3f42" if r_idx % 2 == 0 else "#303437"
        x = 1
        for c_idx, value in enumerate(row):
            draw.rectangle([x, y, x + widths[c_idx], y + row_h], fill=fill, outline="#202326")
            text = ellipsize(draw, value, cell_font, widths[c_idx] - pad_x * 2)
            draw.text((x + pad_x, y + 8), text, fill="#f0f0f0", font=cell_font)
            x += widths[c_idx]
        y += row_h

    img.save(output)


def east_asia(run):
    run.font.name = "맑은 고딕"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")


def add_para(doc, text, style=None, size=10, bold=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    east_asia(run)
    run.font.size = Pt(size)
    run.bold = bold
    return p


def insert_after_anchor(doc, anchor_text, insert_fn):
    for i, p in enumerate(doc.paragraphs):
        if anchor_text in p.text:
            insert_fn(doc, i + 1)
            return True
    return False


def move_paragraph_to(paragraph, parent, index):
    body = parent.element.body
    body.remove(paragraph._p)
    body.insert(index, paragraph._p)


def move_last_element_to(doc, index):
    body = doc.element.body
    element = body[-1]
    body.remove(element)
    body.insert(index, element)


def main():
    IMAGE_DIR.mkdir(parents=True, exist_ok=True)
    tables = [
        (
            "USER",
            ["userID", "userPassword", "userEmail", "userGender", "userType"],
            [
                ["test123", "test123", "test123@doglog.com", "male", "foster"],
                ["user111", "user111", "user111@doglog.com", "female", "adopter"],
            ],
            [150, 180, 260, 160, 160],
        ),
        (
            "APP_USER",
            ["userID", "userPassword", "userEmail", "userGender", "userType"],
            [
                ["user001", "user001", "user001@doglog.com", "female", "foster"],
                ["user01", "user01", "user01@doglog.com", "female", "foster"],
                ["user111", "user111", "user111@doglog.com", "female", "adopter"],
                ["user222", "user222", "user222@doglog.com", "male", "adopter"],
            ],
            [150, 180, 260, 160, 160],
        ),
        (
            "DOG",
            ["id", "name", "breed", "age", "status", "fosterId", "photoData"],
            [
                ["7", "마루", "푸들", "3", "임시보호중", "test123", "/9j/4AAQSkZJRgABAQA..."],
                ["10", "밀크", "말티즈", "13", "임시보호중", "test123", "/9j/4AAQSkZJRgABAQA..."],
                ["12", "호두", "말티푸", "1", "임시보호중", "test123", "/9j/4AAQSkZJRgABAQA..."],
            ],
            [70, 140, 150, 80, 160, 140, 260],
        ),
        (
            "DIARY",
            ["id", "dogName", "dateText", "foodAmount", "poopCount", "content", "photoData", "createdAt", "fosterId"],
            [
                ["10", "마루", "2026.06.11 15:46", "49", "3", "오늘 산책 갔다와서 목욕했어요", "", "2026-06-11 15:46:14", "test123"],
                ["11", "마루", "2026.06.12 04:33", "39", "2", "오늘 애견카페 갔다왔어요", "", "2026-06-12 04:33:19", "test123"],
                ["12", "마루", "2026.06.12 06:17", "46", "1", "test1231443", "/9j/4AAQSkZJRg...", "2026-06-12 06:17:33", "test123"],
            ],
            [60, 110, 170, 120, 120, 300, 190, 190, 120],
        ),
        (
            "CONTACT_MESSAGE",
            ["id", "senderId", "receiverId", "dogName", "title", "content", "createdAt", "replyContent", "repliedAt"],
            [
                ["6", "user111", "test123", "밀크", "밀크 문의", "강아지가 귀엽네요", "2026-06-12 07:27:56", "[NULL]", "[NULL]"],
                ["7", "user222", "test123", "밀크", "밀크 문의", "강아지 입양 문의드립니다", "2026-06-12 08:27:23", "연락처 남겨주세요~", "2026-06-12 08:28:22"],
            ],
            [60, 120, 120, 100, 130, 230, 190, 190, 190],
        ),
    ]

    image_paths = []
    for name, headers, rows, widths in tables:
        out = IMAGE_DIR / f"{name}.png"
        draw_db_table(name, headers, rows, widths, out)
        image_paths.append((name, out))

    doc = Document(SOURCE_DOCX)

    def add_images(doc, insert_index):
        add_para(doc, "DB 테이블 캡처 이미지", "Heading 2", 13, True)
        move_paragraph_to(doc.paragraphs[-1], doc, insert_index)
        insert_index += 1
        add_para(doc, "아래 이미지는 첨부한 DB 테이블 캡처본의 실제 레코드 내용을 보고서에서 바로 볼 수 있도록 재구성한 것이다.", size=9)
        move_paragraph_to(doc.paragraphs[-1], doc, insert_index)
        insert_index += 1
        for name, image_path in image_paths:
            add_para(doc, f"{name} 테이블 캡처", "Heading 3", 11, True)
            move_paragraph_to(doc.paragraphs[-1], doc, insert_index)
            insert_index += 1
            doc.add_picture(str(image_path), width=Inches(6.5))
            move_last_element_to(doc, insert_index)
            insert_index += 1

    if not insert_after_anchor(doc, "DB 테이블 캡처본 기준 실제 저장 레코드", add_images):
        doc.add_page_break()
        add_para(doc, "DB 테이블 캡처 이미지", "Heading 2", 13, True)
        for name, image_path in image_paths:
            add_para(doc, f"{name} 테이블 캡처", "Heading 3", 11, True)
            doc.add_picture(str(image_path), width=Inches(6.5))

    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    main()
